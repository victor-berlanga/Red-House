"""Exporta Markdown técnico a Word; el reporte reutiliza el diseño del primer parcial.

LibreOffice produce el PDF. El catálogo físico conserva su formato independiente.
"""

import argparse
import re
import subprocess
import tempfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import markdown
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import html
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
REPORT_TEMPLATE = (
    ROOT
    / "documentation/historial/primer-parcial/docx/Reporte_Tecnico_Integral_Red_House.docx"
)


def apply_report_design(document):
    """Traslada estilos y portada sin importar capítulos del reporte histórico."""
    template = Document(REPORT_TEMPLATE)
    for name in (
        "Normal",
        "Title",
        "Subtitle",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Caption",
    ):
        original = document.styles[name].element
        original.getparent().replace(original, deepcopy(template.styles[name].element))
    defaults = document.styles.element.find(qn("w:docDefaults"))
    defaults.getparent().replace(
        defaults, deepcopy(template.styles.element.find(qn("w:docDefaults")))
    )
    for section in document.sections:
        for name in (
            "page_width",
            "page_height",
            "top_margin",
            "bottom_margin",
            "left_margin",
            "right_margin",
            "header_distance",
            "footer_distance",
        ):
            setattr(section, name, getattr(template.sections[0], name))
        footer = section.footer._element
        for child in list(footer):
            footer.remove(child)
        for child in template.sections[0].footer._element:
            footer.append(deepcopy(child))

    heading_styles = {
        "Heading 1": "Title",
        "Heading 2": "Heading 1",
        "Heading 3": "Heading 2",
        "Heading 4": "Heading 3",
    }
    for paragraph in document.paragraphs:
        if paragraph.style.name in heading_styles:
            paragraph.style = document.styles[heading_styles[paragraph.style.name]]

    width = (
        document.sections[0].page_width
        - document.sections[0].left_margin
        - document.sections[0].right_margin
    )
    for table in document.tables:
        table.style = "Table Grid"
        table.autofit = False
        borders = deepcopy(template.tables[0]._tbl.tblPr.find(qn("w:tblBorders")))
        previous = table._tbl.tblPr.find(qn("w:tblBorders"))
        if previous is not None:
            table._tbl.tblPr.remove(previous)
        table._tbl.tblPr.append(borders)
        headers = [cell.text for cell in table.rows[0].cells]
        fractions = (
            [0.10, 0.33, 0.39, 0.18]
            if headers == ["Método", "Ruta", "Objetivo", "Códigos declarados"]
            else [1 / len(table.columns)] * len(table.columns)
        )
        for column, fraction in zip(table.columns, fractions):
            column.width = int(width * fraction)
        for index, row in enumerate(table.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell, fraction in zip(row.cells, fractions):
                cell.width = int(width * fraction)
                properties = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "991B1B" if index == 0 else "FFFFFF")
                properties.append(shading)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom"):
                    margin = OxmlElement("w:" + side)
                    margin.set(qn("w:w"), "80")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1
                    for run in paragraph.runs:
                        run.font.name, run.font.size = "Arial", Pt(9)
                        run.font.bold = index == 0
                        run.font.color.rgb = RGBColor.from_string(
                            "FFFFFF" if index == 0 else "1F2933"
                        )

    # La portada termina en el primer salto de página explícito. Se conservan
    # sus párrafos y posiciones originales, incluida la imagen anclada de UDEM.
    cover = []
    for paragraph in template.paragraphs:
        copied = deepcopy(paragraph._p)
        for element in copied.iter():
            for attribute in (qn("r:embed"), qn("r:link"), qn("r:id")):
                identifier = element.get(attribute)
                if not identifier:
                    continue
                relationship = template.part.rels[identifier]
                if relationship.is_external:
                    replacement = document.part.relate_to(
                        relationship.target_ref, relationship.reltype, is_external=True
                    )
                else:
                    replacement, _ = document.part.get_or_add_image(
                        BytesIO(relationship.target_part.blob)
                    )
                element.set(attribute, replacement)
        cover.append(copied)
        if paragraph._p.xpath('.//w:br[@w:type="page"]'):
            break
    else:
        raise ValueError("La plantilla no contiene el salto final de portada.")
    # Evitar identificadores de dibujo repetidos al incorporar la portada.
    next_id = (
        max(
            (
                int(element.get("id"))
                for element in document.element.xpath(".//wp:docPr")
            ),
            default=0,
        )
        + 1
    )
    for paragraph in cover:
        for element in paragraph.xpath(".//wp:docPr"):
            element.set("id", str(next_id))
            next_id += 1
    for paragraph in reversed(cover):
        document.element.body.insert(0, paragraph)
    return len(cover)


def inline_text(paragraph, node):
    """Conserva énfasis y enlaces web, incluidos los títulos de las citas APA."""
    if node.text:
        paragraph.add_run(node.text)
    for child in node:
        if child.tag == "a" and child.get("href", "").startswith(
            ("https://", "http://")
        ):
            link = OxmlElement("w:hyperlink")
            link.set(
                qn("r:id"),
                paragraph.part.relate_to(
                    child.get("href"), RT.HYPERLINK, is_external=True
                ),
            )
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = child.text_content()
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
        else:
            run = paragraph.add_run(child.text_content())
            run.italic = child.tag == "em"
            run.bold = child.tag == "strong"
        if child.tail:
            paragraph.add_run(child.tail)


def insert_contents(document, entries, pages, cover_paragraphs):
    """Índice enlazado; los números se resuelven sobre el PDF ya paginado."""
    # El primer elemento del índice es una sección; el título del reporte debe
    # permanecer al inicio del cuerpo, después del índice.
    anchor = document.paragraphs[cover_paragraphs]._p
    document.paragraphs[cover_paragraphs].paragraph_format.page_break_before = True
    heading = document.add_paragraph("Índice", style="Title")
    anchor.addprevious(heading._p)
    for identifier, (target, depth) in enumerate(entries, 1):
        bookmark = f"report_section_{identifier}"
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(identifier))
        start.set(qn("w:name"), bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(identifier))
        target._p.insert(1 if target._p.pPr is not None else 0, start)
        target._p.append(end)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.left_indent = Inches(0.25 * depth)
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(7.1 - 0.25 * depth), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), bookmark)
        for value in (target.text, "\t", str(pages.get(target.text, "—"))):
            run = OxmlElement("w:r")
            element = OxmlElement("w:tab" if value == "\t" else "w:t")
            if value != "\t":
                element.text = value
            run.append(element)
            link.append(run)
        paragraph._p.append(link)
        anchor.addprevious(paragraph._p)


def export(paths, destination, *, report_design=False, toc_pages=None):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        document.styles[name].font.color.rgb = RGBColor.from_string("8E2434")
    footer = section.footer.paragraphs[0]
    footer.text = "Red House · Segundo parcial · "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    references = {}
    contents = []
    for source_index, source in enumerate(paths):
        root = html.fragment_fromstring(
            markdown.markdown(
                source.read_text(), extensions=["tables", "fenced_code", "md_in_html"]
            ),
            create_parent="div",
        )
        service_number = 0
        for node in root:
            text = node.text_content().strip()
            if node.tag in ("h1", "h2", "h3", "h4"):
                if report_design and source_index == 1:
                    if node.tag == "h1":
                        text = "13. Catálogo de microservicios"
                    elif node.tag == "h2":
                        service_number += 1
                        text = f"13.{service_number}. {text}"
                paragraph = document.add_heading(text, int(node.tag[1]))
                if report_design and source_index > 0 and node.tag == "h1":
                    paragraph.paragraph_format.page_break_before = True
                if report_design and (
                    (source_index == 0 and node.tag == "h2")
                    or (source_index == 1 and node.tag in ("h1", "h2"))
                    or (source_index == 2 and node.tag == "h1")
                ):
                    contents.append(
                        (paragraph, int(source_index == 1 and node.tag == "h2"))
                    )
            elif node.tag == "table":
                rows = node.xpath(".//tr")
                if not rows:
                    continue
                table = document.add_table(rows=0, cols=len(rows[0].xpath("./th|./td")))
                table.style = "Light Shading Accent 1"
                for index, row in enumerate(rows):
                    cells = table.add_row().cells
                    for cell, content in zip(cells, row.xpath("./th|./td")):
                        inline_text(cell.paragraphs[0], content)
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(8)
                    if index == 0:
                        repeat = OxmlElement("w:tblHeader")
                        table.rows[-1]._tr.get_or_add_trPr().append(repeat)
                document.add_paragraph()
            elif node.tag in ("ul", "ol"):
                for item in node.xpath("./li"):
                    paragraph = document.add_paragraph(
                        style="List Bullet" if node.tag == "ul" else "List Number",
                    )
                    inline_text(paragraph, item)
            elif node.tag == "pre":
                paragraph = document.add_paragraph(text)
                for run in paragraph.runs:
                    run.font.name, run.font.size = "Consolas", Pt(8)
            elif node.tag == "p" and node.xpath(".//img"):
                for image in node.xpath(".//img"):
                    path = (
                        (source.parent / image.get("src")).resolve().with_suffix(".png")
                    )
                    if path.exists():
                        with Image.open(path) as bitmap:
                            width, height = bitmap.size
                        scale = min(6.7 / width, 8 / height)
                        document.add_picture(str(path), width=Inches(width * scale))
                        document.add_paragraph(image.get("alt", ""), style="Caption")
                    else:
                        raise FileNotFoundError(path)
            elif text:
                paragraph = document.add_paragraph()
                inline_text(paragraph, node)
                if source.name == "Referencias.md":
                    paragraph.paragraph_format.left_indent = Inches(0.5)
                    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        links = sorted(
            {(a.text_content(), a.get("href")) for a in root.xpath(".//a[@href]")}
        )
        for label, href in links:
            if href.startswith(("http://", "https://", "#")):
                continue
            href = str((source.parent / href).resolve().relative_to(ROOT))
            references.setdefault(href, set()).add(label)
    if references:
        document.add_heading("Documentación y artefactos del proyecto", 2)
        for href, labels in sorted(references.items()):
            document.add_paragraph(
                " / ".join(sorted(labels)) + ": " + href, style="Caption"
            )
    if report_design:
        cover_paragraphs = apply_report_design(document)
        insert_contents(document, contents, toc_pages or {}, cover_paragraphs)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return [paragraph.text for paragraph, _ in contents]


def export_pdf(paths, destination):
    """Repite la paginación hasta estabilizar las páginas del índice enlazado."""
    import pymupdf

    pages = {}
    with tempfile.TemporaryDirectory(prefix="redhouse-report-") as folder:
        temporary = Path(folder)
        for attempt in range(3):
            headings = export(paths, destination, report_design=True, toc_pages=pages)
            subprocess.run(
                [
                    "soffice",
                    f"-env:UserInstallation={(temporary / 'profile').as_uri()}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(temporary),
                    str(destination),
                ],
                check=True,
                timeout=180,
            )
            with pymupdf.open(temporary / destination.with_suffix(".pdf").name) as pdf:
                resolved = {}
                for number, page in enumerate(pdf, 1):
                    for block in page.get_text("dict")["blocks"]:
                        if "lines" not in block:
                            continue
                        text = "".join(
                            span["text"]
                            for line in block["lines"]
                            for span in line["spans"]
                            if span["size"] >= 14.5
                        )
                        for title in headings:
                            if re.sub(r"\s+", "", title) == re.sub(r"\s+", "", text):
                                resolved[title] = number
                if set(resolved) != set(headings):
                    raise ValueError(
                        f"Títulos sin página: {set(headings) - set(resolved)}"
                    )
            if resolved == pages:
                target = (
                    ROOT / "documentation/pdf" / destination.with_suffix(".pdf").name
                )
                target.write_bytes((temporary / target.name).read_bytes())
                print(f"PDF e índice verificados en {attempt + 1} pasadas.")
                return
            pages = resolved
    raise ValueError("La paginación del índice no se estabilizó.")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--physical", action="store_true")
    parser.add_argument(
        "--pdf",
        action="store_true",
        help="Genera PDF e índice con páginas verificadas; requiere LibreOffice y PyMuPDF.",
    )
    args = parser.parse_args()
    sources = (
        [ROOT / "documentation/segundo-parcial/Modelo_fisico.md"]
        if args.physical
        else [
            ROOT / "documentation/markdowns/Reporte_tecnico_segundo_parcial.md",
            ROOT / "documentation/segundo-parcial/Fichas_microservicios.md",
            ROOT / "documentation/segundo-parcial/Referencias.md",
        ]
    )
    name = (
        "Modelo_fisico_PostgreSQL"
        if args.physical
        else "Reporte_tecnico_segundo_parcial"
    )
    if args.pdf and not args.physical:
        export_pdf(sources, ROOT / "documentation/docx" / (name + ".docx"))
        return
    export(
        sources,
        ROOT / "documentation/docx" / (name + ".docx"),
        report_design=not args.physical,
    )
    print(name + ".docx")


if __name__ == "__main__":
    main()

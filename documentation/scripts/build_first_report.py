"""Construye Word y PDF autosuficientes desde Markdown, diagramas y evidencia.

Instalar requirements-report.txt en un entorno documental separado.
No modifica los documentos especializados ni el reporte integral antecedente.
"""
from pathlib import Path
from html import escape
import hashlib
import json
import re
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from markdown_it import MarkdownIt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak,KeepTogether
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'documentation/markdowns/Reporte_Tecnico_del_Primer_Avance.md'
OUT=ROOT/'documentation/docx/Reporte_Tecnico_del_Primer_Avance.docx'
PDF=ROOT/'documentation/pdf/Reporte_Tecnico_del_Primer_Avance.pdf'
EVIDENCE=ROOT/'documentation/evidence/ampliacion-monolito'
md=MarkdownIt('commonmark').enable('table')


def inline_html(children):
    result=[]
    for t in children or []:
        if t.type in ('text','code_inline'):result.append(escape(t.content))
        elif t.type=='strong_open':result.append('<b>')
        elif t.type=='strong_close':result.append('</b>')
        elif t.type=='em_open':result.append('<i>')
        elif t.type=='em_close':result.append('</i>')
        elif t.type in ('softbreak','hardbreak'):result.append('<br/>')
        elif t.type=='link_open':result.append('<link href="'+escape(t.attrGet('href'),quote=True)+'" color="#9f1239">')
        elif t.type=='link_close':result.append('</link>')
    return ''.join(result)


def word_inline(p,children):
    bold=italic=False;href=None
    for t in children or []:
        if t.type=='strong_open':bold=True
        elif t.type=='strong_close':bold=False
        elif t.type=='em_open':italic=True
        elif t.type=='em_close':italic=False
        elif t.type=='link_open':href=t.attrGet('href')
        elif t.type=='link_close':href=None
        elif t.type in ('softbreak','hardbreak'):p.add_run().add_break()
        elif t.type in ('text','code_inline'):
            if href:
                link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(href,RT.HYPERLINK,is_external=True))
                run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'9F1239');props.append(color);run.append(props)
                text=OxmlElement('w:t');text.text=t.content;run.append(text);link.append(run);p._p.append(link)
            else:
                run=p.add_run(t.content);run.bold=bold;run.italic=italic


def results_text():
    file=EVIDENCE/'Tests.xml'
    if not file.exists():raise RuntimeError('Falta el informe final Tests.xml; no se genera un reporte sin resultado verificable.')
    suites=list(ET.parse(file).getroot().iter('testsuite'))
    counts={k:sum(int(s.get(k,0)) for s in suites) for k in ('tests','failures','errors','skipped')}
    seconds=sum(float(s.get('time',0)) for s in suites)
    if counts['failures'] or counts['errors'] or counts['skipped']:
        raise RuntimeError('La suite final contiene fallos u omisiones: revisar antes de publicar el reporte.')
    return f"**Resultado de la suite de esta versión:** {counts['tests']} pruebas aprobadas, 0 fallos, 0 errores y 0 omisiones; {seconds:.2f} segundos. Incluye regresión, migración, integración regional y recorridos de Chrome. Informe: documentation/evidence/ampliacion-monolito/Tests.xml.\n\n",counts


class Report(SimpleDocTemplate):
    def afterFlowable(self,flowable):
        if isinstance(flowable,Paragraph) and hasattr(flowable,'_bookmark'):
            key,level=flowable._bookmark
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(flowable.getPlainText(),key,level=level,closed=False)
            self.notify('TOCEntry',(level,flowable.getPlainText(),self.page,key))


def main():
    result,counts=results_text()
    source=SOURCE.read_text().replace('## 33. Plan y resultados de pruebas\n','## 33. Plan y resultados de pruebas\n\n'+result)
    catalog=json.loads((EVIDENCE/'Physical_schema.json').read_text())
    dictionary=['### Diccionario físico comprobado en PostgreSQL',
        'Las siguientes columnas y restricciones proceden del catálogo real del esquema actualizado. PK: clave primaria; FK: referencia; CHECK: dominio; UNIQUE: unicidad. Los índices operativos completos se conservan en Physical_schema.json.']
    for name in sorted({c['table_name'] for c in catalog['columns']}):
        dictionary += ['#### '+name,'| Columna | Tipo | Admite nulo |','| --- | --- | --- |']
        for c in catalog['columns']:
            if c['table_name']==name:
                size='('+str(c['character_maximum_length'])+')' if c['character_maximum_length'] else ''
                dictionary.append('| '+c['column_name']+' | '+c['data_type']+size+' | '+c['is_nullable']+' |')
        for c in catalog['constraints']:
            if c['table_name']==name:dictionary.append({'p':'PK','f':'FK','c':'CHECK','u':'UNIQUE'}.get(c['type'],c['type'])+': '+c['definition'].replace('|','/')+'.')
        dictionary.append('')
    source=source.replace('## 26. Diseño MongoDB', '\n'.join(dictionary)+'\n\n## 26. Diseño MongoDB')
    headings=re.findall(r'^## (\d+)\. (.+)$',source,re.M)
    assert [int(n) for n,_ in headings]==list(range(1,38))
    tokens=md.parse(source)
    doc=Document();sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
    sec.different_first_page_header_footer=True
    sec.page_width=Inches(A4[0]/72);sec.page_height=Inches(A4[1]/72)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',28),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('9F1239')
    sec.header.paragraphs[0].text='RED HOUSE · REPORTE TÉCNICO DEL PRIMER AVANCE · 2.0'
    sec.header.paragraphs[0].runs[0].font.size=Pt(8)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Equipo 01 · ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    fontroot=Path('/System/Library/Fonts/Supplemental')
    for name,file in [('Body','Arial.ttf'),('Body-Bold','Arial Bold.ttf'),('Body-Italic','Arial Italic.ttf'),('Body-BoldItalic','Arial Bold Italic.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(fontroot/file)))
    pdfmetrics.registerFontFamily('Body',normal='Body',bold='Body-Bold',italic='Body-Italic',boldItalic='Body-BoldItalic')
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle('ReportBody',fontName='Body',fontSize=10,leading=14,spaceAfter=8))
    styles.add(ParagraphStyle('ReportCell',fontName='Body',fontSize=8,leading=11,spaceAfter=2,wordWrap='CJK'))
    styles.add(ParagraphStyle('ReportCaption',fontName='Body-Italic',fontSize=8,leading=11,spaceAfter=12,textColor=colors.HexColor('#475569')))
    for name,size in [('Title',26),('Heading1',17),('Heading2',13)]:
        styles[name].fontName='Body-Bold';styles[name].fontSize=size;styles[name].leading=size+5;styles[name].textColor=colors.HexColor('#9f1239')
        styles[name].keepWithNext=True
    story=[];index=0;images=[];table_count=0;section_no=0;bookmark_no=0
    width=A4[0]-2*.7*72
    while index<len(tokens):
        token=tokens[index]
        if token.type=='heading_open':
            inline=tokens[index+1];level=int(token.tag[1:]);text=inline.content
            if level==1:
                doc.add_paragraph(text,'Title');story.append(Paragraph(escape(text),styles['Title']))
            else:
                if level==2:
                    section_no+=1
                    if section_no==3:
                        doc.add_page_break();doc.add_heading('Índice de contenido',level=1)
                        for number,title in headings[2:]:doc.add_paragraph(number+'. '+title)
                        doc.add_page_break()
                        story.append(PageBreak());story.append(Paragraph('Índice de contenido',styles['Heading1']))
                        toc=TableOfContents();toc.levelStyles=[ParagraphStyle('TOC1',fontName='Body',fontSize=9,leading=13,leftIndent=0,firstLineIndent=0),ParagraphStyle('TOC2',fontName='Body',fontSize=8,leading=11,leftIndent=12)]
                        story.append(toc);story.append(PageBreak())
                    elif section_no in (16,20,23,24,25,26,29,32,33,37):
                        doc.add_page_break();story.append(PageBreak())
                doc.add_heading(text,level=min(level-1,3))
                para=Paragraph(escape(text),styles['Heading1' if level==2 else 'Heading2'])
                para._report_heading=True
                if section_no>=3 and level<=3:
                    bookmark_no+=1;para._bookmark=('section'+str(bookmark_no),0 if level==2 else 1)
                story.append(para)
            index+=3;continue
        if token.type=='paragraph_open':
            inline=tokens[index+1];children=inline.children or []
            picture=next((c for c in children if c.type=='image'),None)
            if picture:
                file=(SOURCE.parent/picture.attrGet('src')).resolve()
                if not file.is_file():raise RuntimeError('Falta figura: '+str(file))
                iw,ih=PILImage.open(file).size
                scale=min(width/iw,510/ih)
                if ih/iw>1.9:scale=min(width/iw,600/ih)
                doc.add_picture(str(file),width=Inches(iw*scale/72))
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
                doc.add_paragraph(picture.content,'Caption')
                figure=[Image(str(file),width=iw*scale,height=ih*scale),Paragraph(escape(picture.content),styles['ReportCaption'])]
                if story and getattr(story[-1],'_report_heading',False):
                    figure.insert(0,story.pop())
                story.append(KeepTogether(figure))
                images.append(str(file.relative_to(ROOT)))
            else:
                p=doc.add_paragraph();word_inline(p,children)
                story.append(Paragraph(inline_html(children),styles['ReportBody']))
            index+=3;continue
        if token.type=='table_open':
            rows=[];current=[];index+=1
            while tokens[index].type!='table_close':
                if tokens[index].type=='tr_open':current=[]
                elif tokens[index].type=='inline':current.append(tokens[index])
                elif tokens[index].type=='tr_close':rows.append(current)
                index+=1
            table_count+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1';table.autofit=False
            for ri,row in enumerate(rows):
                cells=table.rows[0].cells if ri==0 else table.add_row().cells
                for ci,cell in enumerate(row):
                    word_inline(cells[ci].paragraphs[0],cell.children)
                    for p in cells[ci].paragraphs:
                        p.paragraph_format.space_after=Pt(3)
                        for run in p.runs:run.font.size=Pt(8)
                props=table.rows[ri]._tr.get_or_add_trPr()
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');props.append(repeat)
                no_split=OxmlElement('w:cantSplit');props.append(no_split)
            doc.add_paragraph()
            cols=len(rows[0]);colwidths=[width/cols]*cols
            if cols>=4:colwidths=[width*.18]+[width*.82/(cols-1)]*(cols-1)
            data=[[Paragraph(inline_html(cell.children),styles['ReportCell']) for cell in row] for row in rows]
            pt=Table(data,colWidths=colwidths,repeatRows=1,hAlign='LEFT')
            pt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e2e8f0')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#cbd5e1')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            story.extend([pt,Spacer(1,10)])
        index+=1
    OUT.parent.mkdir(parents=True,exist_ok=True);PDF.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    def footer_pdf(canvas,document):
        if document.page<=1:return
        canvas.saveState();canvas.setFont('Body',8);canvas.setFillColor(colors.HexColor('#64748b'))
        canvas.drawString(.7*72,A4[1]-.4*72,'RED HOUSE · REPORTE TÉCNICO DEL PRIMER AVANCE · 2.0')
        canvas.drawRightString(A4[0]-.7*72,.35*72,'Equipo 01 · '+str(document.page));canvas.restoreState()
    report=Report(str(PDF),pagesize=A4,rightMargin=.7*72,leftMargin=.7*72,topMargin=.7*72,bottomMargin=.7*72,
                  title='Reporte Técnico del Primer Avance — Red House',author='Equipo 01')
    report.multiBuild(story,onFirstPage=footer_pdf,onLaterPages=footer_pdf)
    manifest=dict(sections=len(headings),tables=table_count,images=images,tests=counts,source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        artifacts={str(f.relative_to(ROOT)):hashlib.sha256(f.read_bytes()).hexdigest() for f in (OUT,PDF)})
    (EVIDENCE/'Report_build.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps({'result':'PASS','sections':len(headings),'tables':table_count,'images':len(images),'docx':str(OUT),'pdf':str(PDF)},ensure_ascii=False))


if __name__=='__main__':main()
